# -*- coding: utf-8 -*-
"""
生成论文手稿（docx）：python3 paper_figures/src/manuscript.py
需先运行 make_all.py 生成图表与 retrieval_eval.json。
正文中【 】内为需作者补充或核实的内容，在 Word 中以黄色高亮显示。
"""
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

import style
style.setup()
import charts
import tables
from style import OUT

PNG = os.path.join(OUT, "png")
R = json.load(open(os.path.join(OUT, "tables", "retrieval_eval.json"), encoding="utf-8"))
T, O, L = R["template"], R["open"], R["latency_ms"]


def f3(x):
    return f"{x:.3f}"


# ------------------------------------------------------------------ 排版工具
def set_font(run, size=10.5, east="宋体", west="Times New Roman", bold=False, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = west
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), east)
    rf.set(qn("w:ascii"), west)
    rf.set(qn("w:hAnsi"), west)


TOKEN = re.compile(r"(【[^】]*】|\[\d+(?:[,，\-–]\d+)*\])")


def rich(p, text, size=10.5, east="宋体", bold=False, italic=False):
    """【…】 黄色高亮（待补/核实）；[n] 上标引文。"""
    for part in TOKEN.split(text):
        if not part:
            continue
        r = p.add_run(part)
        set_font(r, size, east, bold=bold, italic=italic)
        if part.startswith("【"):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif re.fullmatch(r"\[\d+(?:[,，\-–]\d+)*\]", part):
            r.font.superscript = True


def para(doc, text, size=10.5, east="宋体", bold=False, align="justify", indent=True, before=0, after=0,
         spacing=1.5, italic=False):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    pf = p.paragraph_format
    pf.first_line_indent = Pt(size * 2) if indent else None
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = spacing
    rich(p, text, size, east, bold, italic)
    return p


def heading(doc, text, level):
    size = {1: 12, 2: 10.5, 3: 10.5}[level]
    p = doc.add_heading(level=level)
    p.paragraph_format.space_before = Pt(8 if level == 1 else 4)
    p.paragraph_format.space_after = Pt(4 if level == 1 else 2)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    set_font(r, size, "黑体", bold=(level == 1))
    r.font.color.rgb = None
    return p


def labeled(doc, label, text, size=9, east="宋体", label_east="黑体", italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label)
    set_font(r, size, label_east, bold=True)
    rich(p, text, size, east, italic=italic)
    return p


FIG_N = [0]


def figure(doc, png, zh, en, note=None, width=15.5):
    FIG_N[0] += 1
    n = FIG_N[0]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(os.path.join(PNG, png), width=Cm(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(0)
    r = c.add_run(f"图{n}  {zh}")
    set_font(r, 9, "黑体", bold=True)
    c2 = doc.add_paragraph()
    c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c2.paragraph_format.space_after = Pt(2 if note else 8)
    r = c2.add_run(f"Fig.{n}  {en}")
    set_font(r, 9, "宋体")
    if note:
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p3.paragraph_format.space_after = Pt(8)
        p3.paragraph_format.line_spacing = 1.2
        rich(p3, "注：" + note, 8)
    return n


# 表格编号重映射：tables.py 中的统计表在手稿里重新编号
TABLE_MAP = {}


def remap_add_table(orig):
    def wrapper(doc, title, header, rows, note=None, widths=None, align=None):
        old = title.split(" ")[0]
        new = TABLE_MAP.get(old, old)
        title = title.replace(old, new, 1)
        keep = tables.TDIR
        tables.TDIR = "/tmp/claude-0/ms_tables"          # 手稿重编号的 CSV 不覆盖正式表格
        os.makedirs(tables.TDIR, exist_ok=True)
        try:
            return orig(doc, title, header, rows, note=note, widths=widths, align=align)
        finally:
            tables.TDIR = keep
    return wrapper


tables.add_table = remap_add_table(tables.add_table)


# ------------------------------------------------------------------ 正文
def build():
    topo = charts.fig_topology()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    st.font.size = Pt(10.5)
    for lvl in (1, 2, 3):
        hs = doc.styles[f"Heading {lvl}"]
        hs.font.color.rgb = None

    m5, m3, m4, m2, m1 = T["M5"], T["M3"], T["M4"], T["M2"], T["M1"]
    o5, o4, o3 = O["M5"]["120"], O["M4"]["120"], O["M3"]["120"]

    # ---------------- 题名页 ----------------
    para(doc, "孙光荣国医大师中和学术思想专题知识图谱构建与可溯源智能问答研究", 18, "黑体", True, "center",
         False, 12, 8, 1.3)
    para(doc, "【作者1】¹，【作者2】¹，【作者3】²，【通信作者】¹*", 12, "仿宋", align="center", indent=False, after=4)
    para(doc, "（1. 【单位1，城市 邮编】；2. 【单位2，城市 邮编】）", 9, align="center", indent=False, after=10)

    para(doc, "摘  要", 10.5, "黑体", True, "left", False, after=2)
    abstract = [
        ("目的", "构建面向国医大师孙光荣“中和”学术思想的专题知识图谱，并在此基础上实现回答可回溯至原文的智能问答，"
                 "为名老中医学术思想的数字化传承提供方法参考。"),
        ("方法", "以孙光荣著述及其编纂文献 8 部为语料，按段落切分为 1 734 个可溯源文本块；设计包含中和思想、辨证论治、方药配伍、"
                 "医案诊次、养生康复、传承溯源 6 层 31 类概念与 44 类关系的本体，抽取实体与关系，并为每条关系绑定原文佐证、"
                 "出处编号与 5 类知识来源标签；以本体一致性校验与抽样人工审校控制质量；运用网络拓扑分析、类型层链路流分析、"
                 "Louvain 社团划分与关联规则挖掘分析图谱；提出融合实体—佐证双路 BM25 召回、一跳图扩展证据加权与关系类型轮转采样的"
                 "可溯源图检索增强生成（GraphRAG）方法，构建 234 道模板问答题与 30 道开放问题，与 4 种检索方案比较。"),
        ("结果", f"图谱含实体 19 343 个、关系 25 713 条，带原文佐证的关系占 99.98%；v1、v2 版抽样严格精度分别为 84.9%（n=119）"
                 f"和 81.7%（n=60），801 条（3.1%）关系被标记待审。孙光荣原创关系 9 186 条（35.7%），中和思想层与医案诊次层的"
                 f"孙光荣原创实体分别占 60% 与 91%。专题子图呈现“经典渊源—中和思想—治则—治法—药组”知识链，涵盖治则 27 条；"
                 f"290 首医案处方中生黄芪、紫丹参出现率分别为 85.2% 和 84.5%，含三联药组的处方中 86.9% 使用“参—生黄芪—紫丹参”"
                 f"益气活血基础药组。所提方法 Hit@20 为 {f3(m5['hit@20'])}、Recall@20 为 {f3(m5['recall@20'])}，均为各方案最高；"
                 f"Hit@10 为 {f3(m5['hit@10'])}，略低于实体一跳检索（{f3(m3['hit@10'])}）；开放问题中单一关系类型最大占比降至 "
                 f"{f3(o5['max_share'])}，实质佐证占比 {f3(o5['substantive'])}，单次检索中位时延 {L['M5']['median']:.1f} ms。"),
        ("结论", "专题知识图谱能够结构化表达孙光荣中和学术思想的源流、内涵与临证落实路径；可溯源 GraphRAG 在保持较高召回的同时"
                 "提高了检索证据的多样性与可核查性，可为名老中医学术思想传承提供可复用的技术路径。"),
    ]
    for k, v in abstract:
        labeled(doc, f"{k}  ", v)
    labeled(doc, "关键词：", "孙光荣；国医大师；中和学术思想；知识图谱；检索增强生成；可溯源问答；三联药组")
    labeled(doc, "中图分类号：", "R2-03；TP391    文献标志码：A    文章编号：【 】")
    labeled(doc, "基金项目：", "【项目名称（编号）】")
    labeled(doc, "作者简介：", "【姓名（出生年—），性别，学历，职称，研究方向：……；E-mail：……】")

    para(doc, "Construction of a Thematic Knowledge Graph of Zhonghe Academic Thought of TCM Master Sun Guangrong "
              "and Traceable Intelligent Question Answering", 12, bold=True, align="center", indent=False,
         before=12, after=6, spacing=1.2)
    para(doc, "【AUTHOR1】¹, 【AUTHOR2】¹, 【AUTHOR3】², 【CORRESPONDING AUTHOR】¹*", 10, align="center", indent=False)
    para(doc, "(1. 【Affiliation 1, City Postcode, China】; 2. 【Affiliation 2】)", 9, align="center", indent=False, after=6)
    en = [
        ("Objective: ", "To construct a thematic knowledge graph (KG) of the Zhonghe (centred harmony) academic thought of "
                        "National TCM Master Sun Guangrong and to build an intelligent question-answering system whose answers "
                        "can be traced back to the original texts."),
        ("Methods: ", "Eight monographs authored or compiled by Sun Guangrong were split into 1 734 traceable paragraph chunks. "
                      "An ontology of six layers (Zhonghe thought, pattern differentiation and treatment, formula and herb "
                      "compatibility, case records, health preservation and rehabilitation, and academic lineage) with 31 concept "
                      "types and 44 relation types was designed. Each extracted relation was bound to its textual evidence, chunk "
                      "identifier and one of five knowledge-source labels. Quality was controlled by ontology consistency checking "
                      "and sampled manual auditing. Network topology, type-level chain flows, Louvain communities and association "
                      "rules were analysed. A traceable graph retrieval-augmented generation (GraphRAG) method combining dual-path "
                      "BM25 recall over entities and evidence, evidence-weighted one-hop expansion and relation-type round-robin "
                      "sampling was proposed and compared with four retrieval baselines on 234 template questions and 30 open "
                      "questions."),
        ("Results: ", f"The KG contains 19 343 entities and 25 713 relations, 99.98% of which carry textual evidence. Strict "
                      f"precision of sampled auditing was 84.9% (n=119) for v1 and 81.7% (n=60) for v2; 801 relations (3.1%) "
                      f"were flagged for review. Sun's original assertions account for 9 186 relations (35.7%), and for 60% and "
                      f"91% of entities in the Zhonghe-thought and case-record layers. The thematic subgraph presents a chain of "
                      f"classical origins - Zhonghe concepts - treatment principles - methods - herb triplets, covering 27 "
                      f"principles. In 290 prescriptions, Shenghuangqi and Zidanshen appeared in 85.2% and 84.5%, and 86.9% of "
                      f"triplet-based prescriptions used a ginseng-astragalus-salvia triplet. The proposed method achieved the "
                      f"highest Hit@20 ({f3(m5['hit@20'])}) and Recall@20 ({f3(m5['recall@20'])}); its Hit@10 "
                      f"({f3(m5['hit@10'])}) was slightly lower than entity one-hop retrieval ({f3(m3['hit@10'])}). On open "
                      f"questions the maximum single-relation-type share fell to {f3(o5['max_share'])}, the substantive-evidence "
                      f"ratio was {f3(o5['substantive'])}, and median retrieval latency was {L['M5']['median']:.1f} ms."),
        ("Conclusion: ", "The thematic KG structurally represents the origins, connotations and clinical realisation of Sun "
                         "Guangrong's Zhonghe thought, and traceable GraphRAG improves the diversity and verifiability of "
                         "retrieved evidence while keeping high recall, offering a reusable path for inheriting the academic "
                         "thought of eminent TCM physicians."),
        ("Keywords: ", "Sun Guangrong; National TCM Master; Zhonghe academic thought; knowledge graph; retrieval-augmented "
                       "generation; traceable question answering; herb triplet"),
    ]
    for k, v in en:
        labeled(doc, k, v, east="Times New Roman", label_east="Times New Roman")

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)

    # ---------------- 引言 ----------------
    P = lambda t: para(doc, t)
    P("名老中医学术思想与临证经验是中医药学术传承的核心载体。孙光荣教授为第二届国医大师，长期致力于中医临床、文献与教育，"
      "首倡“中和”学术思想并形成“中和医派”，提出“中和思想—中和辨证—中和组方”的诊疗思路，以及以“三联药组”为核心的组方用药特色"
      "[1-2]。其学术内容分散于专著、医案、心法要诀与选编文献之中，既有孙老本人的论断，也大量引述经典、选编他人经验，"
      "传统的阅读与归纳方式难以完整呈现其思想源流与临证落实之间的结构关系，也难以区分“谁说的”“出自哪里”。")
    P("知识图谱以“实体—关系—实体”三元组组织知识，便于表达概念层级与推理链路[3-4]，近年已用于中医药理论、方剂与名老中医经验的"
      "结构化表示【此处补充 2–4 篇近年中医药知识图谱/名老中医知识图谱中文核心期刊文献】。与此同时，大语言模型为中医知识问答提供了"
      "新途径，但其“幻觉”问题在医学场景中尤为突出[5-6]。检索增强生成（retrieval-augmented generation，RAG）通过先检索外部证据"
      "再生成答案来缓解这一问题[7-8]，基于图结构的检索增强（GraphRAG）进一步利用实体关系组织证据[9-10]。然而，现有中医知识图谱"
      "多面向通用知识，较少围绕某一位名老中医的学术思想做专题化建模；常规文本块 RAG 切断了知识之间的结构联系，"
      "也难以把答案中的每一条论断回溯到原文，更无法区分名老中医本人的观点与其选编、引用的内容。")
    P("针对上述问题，本研究以孙光荣中和学术思想为主题：①设计六层专题本体，为每条关系绑定原文佐证、出处编号与知识来源标签，"
      "构建可溯源的专题知识图谱；②从规模质量、网络拓扑、辨证论治链路、中和思想专题子图与三联药组用药规律等角度分析图谱；"
      "③提出面向该图谱的可溯源 GraphRAG 问答方法，并以可复现的检索实验检验其效果，为名老中医学术思想的数字化传承提供参考。"
      "研究技术路线见图1。")
    figure(doc, "图1_研究技术路线.png", "研究技术路线", "Technical roadmap of the study",
           "红框（D1—D3）为中和学术思想核心文献；抽取融合层的具体实现见 1.3。")

    # ---------------- 1 资料与方法 ----------------
    heading(doc, "1  资料与方法", 1)
    heading(doc, "1.1  数据来源", 2)
    P("以孙光荣本人著述及其主编、选编的 8 部文献为语料（附表 S2）：《国医大师孙光荣中和思想与临证经验集萃》（D1）、"
      "《孙光荣医案解读》（D2）、《医道中和——国医大师孙光荣临证心法要诀》（D3）、《中华经典养生名言录》（D4）、"
      "《中医康复研究》（D5）、《炎症的中医辨治》（D6）、《’91中医新方妙术》（D7）、《精方妙药与美容》（D8）"
      "【请补充各书作者/主编、出版社与出版年】。其中 D1—D3 集中阐述中和学术思想与临证经验，为本研究的核心文献。"
      "各文献按段落切分为文本块，并以“文献编号#文本块编号”（如 D1#c0072）唯一标识，共 1 734 个，作为全部知识的最小溯源单元。")
    heading(doc, "1.2  专题本体设计", 2)
    P("参照中医“理法方药”的诊疗逻辑与孙光荣“中和思想—中和辨证—中和组方”的学术框架，自顶向下设计六层本体（图2，表1）："
      "①中和思想层，含中和思想、学术观点、心法要诀、名言；②辨证论治层，含症状、舌象、脉象、疾病、证候、病因、病机、治则、治法；"
      "③方药配伍层，含方剂、处方、新方、三联药组、药对、中药、外用制剂、食疗方；④医案诊次层，含医案、诊次、患者特征；"
      "⑤养生康复层，含养生法、康复技术、美容适应证；⑥传承溯源层，含经典文献、医家、学派、文献出处。共定义 31 类概念与 44 类关系"
      "（附表 S1），每类关系规定定义域与值域。其中“思想含治则”“治则指导治法”“治则见于药组”与“药组功效”等关系贯通中和思想层、"
      "辨证论治层与方药配伍层，用于刻画中和思想的临证落实路径。")
    P("为支撑可溯源问答，为每条关系设置 4 项溯源属性：原文佐证（引文 q）、出处编号（c）、所属文献（d）与知识来源（r）。知识来源分"
      "5 类：孙光荣原创（孙老本人论断）、孙光荣编纂（孙老选编他人内容）、经典引文、他人临床报道与中医通识。实体另设中英文名称与异名。")
    figure(doc, "图2_中和学术思想专题本体模型.png", "孙光荣中和学术思想专题知识图谱本体模型",
           "Ontology of the thematic knowledge graph of Sun Guangrong's Zhonghe academic thought",
           "节点为概念类型，数字为实体数；边为关系类型，数字为关系实例数，线宽∝lg(实例数)；红色为“中和思想—治则—治法—药组”贯通链路，"
           "灰色虚线为跨层溯源与关联关系；“食疗方”已定义、暂无实例。")
    TABLE_MAP.update({"表1": "表1"})
    tables.table1(doc)

    heading(doc, "1.3  知识抽取与融合", 2)
    P("以文本块为单位，依据本体进行实体识别、类型判定与关系抽取【请据实说明抽取方式，例如：采用××大语言模型在本体约束提示下零样本/少样本抽取，"
      "或人工标注与模型抽取相结合；并说明人工校对的轮次与人员】。每条关系保留支持该关系的原文片段作为佐证，并记录其出处文本块与知识来源。"
      "随后进行实体对齐：同一概念的不同表述归并为异名（共 1 889 条，涉及 1 470 个实体），跨文献出现的同一概念保留各自出处"
      "（1 757 个实体见于 2 部及以上文献）。")
    heading(doc, "1.4  质量控制", 2)
    P("采用两级质量控制。①本体一致性校验：对每条关系检查其头、尾实体类型是否符合该关系的定义域与值域，不符者在图谱中标记为“待审”，"
      "在问答中降低权重并提示用户。②抽样人工审校：对 v1、v2 两版图谱分别随机抽取 119 条与 60 条关系，由【审校人员资质与人数】依据原文"
      "判定其是否完全正确（严格精度），并以 Wilson 区间[16]估计 95% 置信区间。")
    heading(doc, "1.5  图谱分析方法", 2)
    P("①规模与构成：统计各层实体、各类关系及知识来源分布。②网络拓扑：将图谱视为无向图，计算度分布、连通分量，并以离散幂律最大似然"
      "估计（以 KS 距离最小确定 x_min）[13]检验度分布的重尾特征；全图可视化采用 ForceAtlas2 布局[12]。③辨证论治链路：将实体类型归入"
      "“四诊信息—病证—病因病机—治则治法—方剂—药组—中药”7 个环节，统计环节间关系实例数并绘制流图。④中和思想专题子图：以 17 个"
      "“中和思想”实体为中心，取其渊源、治则及治则所指导的治法、药组构成子图。⑤用药规律：从 D2 医案中提取 290 首处方的药物组成"
      "（处方以三联药组编码者按药组名拆分），统计药物与三联药组频次，以 Louvain 算法[14]对高频药物共现网络进行社团划分，"
      "并按支持度≥0.08、置信度≥0.80、提升度≥1.30 挖掘关联规则[15]。")

    heading(doc, "1.6  可溯源 GraphRAG 智能问答方法", 2)
    P("问答系统由检索与生成两部分组成（图3），检索部分以 numpy 实现，无需向量数据库与 GPU。")
    P("（1）查询解析与双路召回。中文问题切分为字符二元组，长度≤6 的短词另作整体词元。建立两个 BM25 索引[11]：实体索引覆盖实体名称、"
      "异名、英文名与类型（19 343 个实体）；佐证索引覆盖每条关系的原文引文及其关系名与两端实体名（25 709 条）。中医术语高度专名化，"
      "佐证原文往往直接包含答案，两路互补。")
    P("（2）种子融合。实体得分为实体索引 BM25 得分，加上佐证索引前 60 条命中引文回流至两端实体的分值（0.55w，w 为引文得分），"
      "若实体名或异名完整出现在问题中再加 4+0.9|s|（|s| 为名称长度）。同名实体去重后取前 35 个作为种子。多轮对话中，"
      "上一轮得分最高的 3 个种子实体并入下一轮查询，用于指代消解。")
    P("（3）一跳图扩展与证据加权。对每个种子取其一跳邻接关系，按式（1）打分：")
    para(doc, "s(e) = ŝ_seed + q(e) + 0.55ρ(r_e) + 0.30·min(ŝ_o, 1) − 0.12·min(ln(1+d_o)/6, 1) − 0.25·𝟙[待审]        （1）",
         10, align="center", indent=False)
    P("式中 ŝ_seed、ŝ_o 分别为种子与对端实体相对最高种子的归一化得分；q(e) 为证据质量项，引文≥10 字加 0.50、4—9 字加 0.16，"
      "引文仅重复实体名者减 0.40；ρ(r) 为知识来源权重（孙光荣原创 1.00、孙光荣编纂 0.72、经典引文 0.60、他人临床报道 0.50、"
      "中医通识 0.42），体现“优先采信孙老本人论断”；d_o 为对端实体的度，用于抑制枢纽实体；待审关系减 0.25。每个种子每类关系"
      "至多入选 3 条，佐证索引直接命中的关系额外加 1.2 以确保保留。")
    P("（4）关系类型轮转采样。枢纽实体（如“失眠”）常挂接数十条同类关系，按分数截断会使结果被同质三元组占满。本文将候选关系按关系类型"
      "分桶，按各桶最高分排序后轮转取样，单一类型至多占 1/4，总数至多 120 条；关系类型本来较少时自动退化为按分数排序。")
    P("（5）可引用上下文与生成约束。入选关系依次编号为 [R1]、[R2]……，与原文佐证、出处编号、文献名、知识来源一并组装为上下文"
      "（上限 9 900 字，实体与关系分账）。系统提示要求模型：只依据图谱证据作答，证据中没有的明确说明“图谱中未见记载”；引用时标注 [Rk] "
      "并给出出处编号；区分孙光荣原创与编纂、经典引文、他人报道；引用待审关系须注明。生成端通过 OpenAI 兼容接口调用"
      "【实际使用的大模型名称与版本】，界面同步展示检索子图、三元组与原文，供用户逐条核对（图4）。")
    figure(doc, "图10_可溯源GraphRAG问答框架.png", "可溯源 GraphRAG 智能问答框架",
           "Framework of traceable graph retrieval-augmented question answering",
           "虚线为多轮对话的指代消解回路；全部参数取自开源实现。")
    figure(doc, "图11_可溯源证据链示例.png", "可溯源证据链示例",
           "An example of the traceable evidence chain",
           "问题为“孙光荣‘中和组方’的基本原则是什么？”；[Rk] 为实际检索排序编号，仅展示与种子“中和组方”直接相连的关系。")

    heading(doc, "1.7  检索实验设计", 2)
    P(f"（1）模板问答集。从知识来源为孙光荣原创或编纂的关系中，按“思想含治则”“治则指导治法”“明机立法”“求因明机”“审证求因”"
      f"“立法组方”“药组功效”“主治”“见症”“渊源”“要诀所述”“引经据典”12 类关系，各随机抽取不超过 20 个不同主语（随机种子 2026），"
      f"以“主语+关系”模板生成问题，如“气滞血瘀应如何立法治疗？”；标准答案为图谱中同主语、同关系类型的全部三元组。共 {R['n_template']} 题。")
    P(f"（2）开放问题集。围绕中和思想内涵与渊源、中和组方、三联药组、常见病辨治、养生与学术传承等撰写 {R['n_open']} 个开放问题，"
      "无标准答案，用于考察检索结果的侧面覆盖与证据质量。")
    P("（3）对比方案。M1 文本块 BM25：以出处编号聚合原文佐证为文本块，块级 BM25 检索后展开其中关系，模拟常规文本 RAG；"
      "M2 佐证句 BM25：直接对每条关系的原文佐证检索；M3 实体+一跳：实体召回后取全部一跳关系，仅按种子得分排序；"
      "M4 本文方法去除轮转采样；M5 本文完整方法。各方案均输出至多 120 条关系。")
    P("（4）评价指标。模板问答集采用 Hit@k（前 k 条中含标准答案的问题比例）、Recall@20 与平均倒数排名（MRR）；"
      "开放问题集采用前 120 条结果的关系类型数、归一化类型熵、单一关系类型最大占比、实质佐证占比（引文≥10 字）与孙光荣原创占比；"
      "另记录单次检索时延。实验仅评测检索层，不调用大模型，结果完全可复现。")

    # ---------------- 2 结果 ----------------
    heading(doc, "2  结果", 1)
    heading(doc, "2.1  图谱规模、构成与质量", 2)
    TABLE_MAP.update({"表4": "表2"})
    P("所构建图谱包含实体 19 343 个、关系 25 713 条（表2），30 类概念已实例化。辨证论治层实体最多（9 025 个），其次为方药配伍层"
      "（4 656 个）、传承溯源层（2 041 个）、养生康复层（1 699 个）、中和思想层（1 231 个）与医案诊次层（691 个）（表1）。"
      "关系以“见症”（3 244 条）、“主治”（1 859 条）为主，“思想含治则”（65 条）、“治则指导治法”（286 条）等中和贯通关系数量虽少，"
      "但处于语义层级的顶端（附图 S1）。25 709 条关系（99.98%）带有原文佐证，全部关系均有出处编号。")
    P("从知识来源看（图5），孙光荣原创关系 9 186 条（35.7%），中医通识 9 053 条（35.2%），他人临床报道 4 200 条（16.3%），"
      "经典引文 2 392 条（9.3%），孙光荣编纂 882 条（3.4%）。孙光荣原创关系集中于 D1—D3，而 D5—D8 以他人报道、经典引文与通识为主；"
      "中和思想层 60%、医案诊次层 91% 的实体为孙光荣原创，表明专题图谱能够在核心层把孙老本人论断与其编纂、引用内容区分开。")
    tables.table7(doc, topo)
    figure(doc, "图4_知识来源与文献分布.png", "知识来源与文献分布",
           "Distribution of knowledge sources across documents and ontology layers",
           "A. 各文献关系实例按知识来源堆叠；B. 各本体层实体的知识来源构成（%）。")
    P("质量方面，v1、v2 版抽样严格精度分别为 84.9%（n=119，95% CI 77.3%—90.2%）与 81.7%（n=60，95% CI 70.1%—89.4%），"
      "两者置信区间重叠，差异不具统计学意义。本体一致性校验共标记待审关系 801 条（3.1%），以“治则指导治法”（37.4%）、"
      "“以方为基”（16.9%）与“辨证为”（14.2%）最高（附图 S2），提示治则与治法的边界是后续人工校对的重点。")

    heading(doc, "2.2  网络拓扑特征", 2)
    P(f"图谱平均度 {topo['mean_deg']:.2f}，{topo['leaf']:.1%} 的实体为度为 1 的叶节点；共 {topo['comps']:,} 个连通分量，"
      f"最大连通分量含 {topo['giant']:,} 个实体（{topo['giant']/19343:.1%}）。度分布呈重尾特征，尾部（k≥{topo['xmin']}）"
      f"幂律拟合指数 α={topo['alpha']:.2f}（KS={topo['ks']:.3f}）。度最高的枢纽实体为“孙光荣”（463）、“中风”（241）、“生甘草”（143）"
      f"与三联药组“生晒参+生黄芪+紫丹参”（141）（附图 S3），反映出图谱以孙光荣为学术中心、以中风等优势病种和核心药组为临床枢纽的结构。"
      .replace(",", " "))

    heading(doc, "2.3  辨证论治链路", 2)
    P("将实体类型按诊疗环节排列后（图6），主链路依次为：症状→证候 1 466 条、舌象与脉象→证候 627 条，证候→病机 341 条、"
      "疾病与证候→病因 615 条，证候→治法 537 条、病机→治法 211 条，治法→方剂 351 条、治则→方剂 302 条，处方→三联药组 1 141 条，"
      "三联药组→中药 872 条。值得注意的是，治法与三联药组之间存在 377 条“药组功效”关系，构成“以法统药组”的独立通路，"
      "这正是孙光荣以三联药组替代单味药行使君臣佐使职能这一组方特色[2]在图谱结构上的体现。")
    figure(doc, "图7_辨证论治链路流图.png", "“症—证—因—机—法—方—药”辨证论治链路流图",
           "Flow of the symptom-syndrome-etiology-pathomechanism-method-formula-herb chain",
           "带宽∝类型间关系实例数，仅显示≥20 条的流量；同一环节内部关系未绘出；红色为经“治则”或“三联药组”的中和特色链路。")

    heading(doc, "2.4  中和学术思想专题子图", 2)
    P("以 17 个“中和思想”实体为中心抽取专题子图（图7，表3）。在渊源上，“中和学术思想”上溯《黄帝内经》《难经》《中藏经》《伤寒论》"
      "《金匮要略》，并引《中庸》《说文》《虞书·大禹谟》释“中”“和”之义，融合朱丹溪、李东垣两家之长；“中和观”引《周易》《国语》"
      "《论语》《道德经》等论述“和实生物”“和为贵”。在内涵上，子图共含治则 27 条，如“扶正祛邪益中和”“护正防邪固中和”"
      "“存正抗邪达中和”“调气血、平升降、衡出入”“清、平、轻、巧、灵”；“中和组方”下含“遵经方之旨，不泥经方用药”"
      "“谨守病机，以平为期”“中病即止，不滥伐无过”“从顺其宜，病人乐于接受”4 条基本原则。在临证落实上，治则经“治则指导治法”"
      "落到“衡出入”“调气血”“护脾胃”“温阳”“滋阴”等治法，并经“治则见于药组”落到“参、芪、紫丹参”药组与具体医案处方，"
      "形成“渊源—思想—治则—治法—药组”的完整知识链，使中和思想由抽象理念变为可查询、可追溯的结构化知识。")
    figure(doc, "图8_中和学术思想专题子图.png", "孙光荣中和学术思想专题子图",
           "Thematic subgraph of Sun Guangrong's Zhonghe academic thought",
           "自左至右为学术渊源、中和思想、治则、治法与药组；同名经典已合并，“孙光荣”节点及心法要诀、名言、医案等表述类节点从略（见表3）。",
           width=15.0)
    TABLE_MAP.update({"表5": "表3"})
    tables.table4(doc)

    heading(doc, "2.5  三联药组用药规律", 2)
    P("290 首医案处方共用药 312 味，平均每方 12.8 味。生黄芪与紫丹参分别见于 85.2% 和 84.5% 的处方，构成核心药对；"
      "274 首以三联药组编码的处方中，238 首（86.9%）含“某参+生黄芪+紫丹参”药组，其中“生晒参+生黄芪+紫丹参”104 首（35.9%）、"
      "“西洋参+生黄芪+紫丹参”74 首（25.5%）、“西党参+生黄芪+紫丹参”31 首、“太子参+生黄芪+紫丹参”18 首（图8A）。"
      "孙光荣依据气虚程度与寒热偏性在生晒参、西洋参、党参、太子参之间择用，而以黄芪、丹参益气活血为恒定基础，"
      "体现“调气血”的中和用药主线【此句为中医学解释，请作者结合临床核实】。")
    P("去除核心药对后，前 32 味高频药的共现网络划分为 4 个社团（图8B）：社团Ⅰ以生甘草、生晒参、制首乌、法半夏、广陈皮、明天麻、"
      "石菖蒲等为代表；社团Ⅱ以云茯神、炒枣仁、西洋参、龙眼肉、大红枣、麦门冬为代表；社团Ⅲ以川杜仲、阿胶珠、全当归、益母草、制香附、"
      "延胡索为代表；社团Ⅳ以山慈菇、蒲公英、白花蛇舌草、猫爪草、生薏米为代表【各社团的功效归纳（如养心安神、养血调经、"
      "清热解毒散结）请作者审定】。关联规则（表4）显示，炒枣仁→云茯神（置信度 0.979，提升度 2.81）、明天麻→制首乌（0.951，4.93）、"
      "法半夏→广陈皮（0.927，6.11）、猫爪草→山慈菇（0.943，3.55）、制远志→石菖蒲（0.968，7.20）等规则支持度高，"
      "与图谱中“云茯神+炒枣仁+龙眼肉”“制首乌+明天麻+蔓荆子”“山慈菇+猫爪草+天葵子”等三联药组相互印证。")
    figure(doc, "图9_三联药组与核心用药网络.png", "三联药组与核心用药共现网络",
           "Herb triplets and core herb co-occurrence network in Sun Guangrong's prescriptions",
           "A. 出现频次前 15 的三联药组，括号内为占全部处方比例；B. 前 32 味高频药共现网络（Louvain 社团划分，节点大小∝频次，"
           "连线为共现≥12 首且提升度>1.2），核心药对置于圆心。")
    TABLE_MAP.update({"表6": "表4"})
    tables.table5(doc)

    heading(doc, "2.6  可溯源问答与检索实验", 2)
    P(f"以“孙光荣‘中和组方’的基本原则是什么？”为例（图4），系统召回 35 个种子实体、{111} 条关系进入上下文，其中与种子“中和组方”"
      "直接相连的关系全部可回溯至同一文本块 D1#c0072，答案中的每条 [Rk] 引用均可定位到原文。索引构建约 2 s，单次检索中位时延 "
      f"{L['M5']['median']:.1f} ms（第 95 百分位 {L['M5']['p95']:.1f} ms）。")
    P(f"检索实验结果见表5与图9。在模板问答集上，本文方法（M5）的 Hit@20 为 {f3(m5['hit@20'])}、Recall@20 为 {f3(m5['recall@20'])}、"
      f"Hit@50 为 {f3(m5['hit@50'])}，均为各方案最高；但其 Hit@1、Hit@10 与 MRR（{f3(m5['hit@1'])}、{f3(m5['hit@10'])}、"
      f"{f3(m5['mrr'])}）低于实体+一跳方案（M3：{f3(m3['hit@1'])}、{f3(m3['hit@10'])}、{f3(m3['mrr'])}）。原因在于模板问题均显式"
      f"给出主语实体，实体中心的检索天然占优；而本文方法为提升证据质量与侧面多样性，对证据薄弱、重复的关系降权，"
      f"损失了少量头部排序精度。常规文本块检索（M1）的 Hit@10 仅 {f3(m1['hit@10'])}，说明以文本块为单位检索难以精准定位具体知识。"
      f"M4 与 M5 在模板问答集上排序结果一致，表明轮转采样不影响针对性问题的命中。")
    P(f"在开放问题集上，轮转采样的作用得以体现：单一关系类型最大占比由 M3 的 {f3(o3['max_share'])}、M4 的 {f3(o4['max_share'])} "
      f"降至 M5 的 {f3(o5['max_share'])}，归一化类型熵升至 {f3(o5['entropy'])}，平均覆盖 {o5['types']:.1f} 类关系；"
      f"实质佐证占比为 {f3(o5['substantive'])}，高于 M1—M3。M1 虽覆盖关系类型最多（{O['M1']['120']['types']:.1f} 类），"
      "但其命中率低，说明“侧面多”并不等同于“相关”。")
    TABLE_MAP.update({"表7": "表5"})
    tables.table6(doc)
    figure(doc, "图12_检索实验结果.png", "检索实验结果", "Results of retrieval experiments",
           f"A. 模板问答集（n={R['n_template']}）Hit@k；B. 分关系类型 Hit@10；C. 开放问题集（n={R['n_open']}）单一关系类型最大占比"
           "（越低越均衡）；D. 开放问题集关系类型归一化熵、实质佐证占比与孙光荣原创占比。M4、M5 在模板问答集上结果一致。")

    # ---------------- 3 讨论 ----------------
    heading(doc, "3  讨论", 1)
    heading(doc, "3.1  专题化建模使中和思想“可结构化、可追问”", 2)
    P("与面向通用中医知识的图谱不同，本研究以一位国医大师的学术思想为主题组织本体，并单设“中和思想层”及贯通至方药层的关系。"
      "结果表明，这种专题化设计能够把分散在多部著作中的论述连接为“渊源—思想—治则—治法—药组”的知识链：读者既可以自上而下追问"
      "某一中和理念如何落实到治法与药组，也可以自下而上追问某一三联药组体现了何种治则与思想。链路流图中 377 条“药组功效”关系"
      "与医案处方中 86.9% 的“参—芪—丹”药组使用率相互印证，从数据层面支持了“三联药组”是孙光荣中和组方核心载体的认识[2]。")
    heading(doc, "3.2  “可溯源”是名老中医知识问答的必要条件", 2)
    P("名老中医学术传承最忌“张冠李戴”。本研究在关系层面同时记录原文佐证、出处编号与知识来源，并在检索打分中赋予孙光荣原创更高权重、"
      "在生成约束中要求区分原创与编纂，使问答系统能够回答“这是谁的观点、出自哪一段”。在证据链示例中，与问题直接相关的关系全部指向"
      "同一文本块，用户可据此核对原文。与需要向量库与 GPU 的方案相比，本方法的检索完全基于稀疏索引与图结构，部署门槛低，"
      "适合在基层与教学场景推广。")
    heading(doc, "3.3  局限与展望", 2)
    P("本研究存在以下局限：①知识抽取的抽样严格精度约 82%—85%，样本量较小、置信区间较宽，约 3.1% 的关系未通过本体校验，"
      "仍需扩大人工审校；②模板问答集由图谱自动生成且显式点名主语实体，不能完全代表真实用户口语化、隐含式的提问；"
      "③本研究仅评测了检索层，大模型最终答案的准确性、完整性、引用正确率与安全性尚需中医专家盲评【若已完成专家评价，请在此补充"
      "方法与结果；若未完成，保留本句作为局限】；④每个种子每类关系至多取 3 条的配额在兼顾多样性的同时可能截断同类信息，"
      "如证据链示例中“中和组方”的第 4 条基本原则“从顺其宜，病人乐于接受”未进入上下文，后续拟引入按问题意图自适应调整配额的策略；"
      "⑤语料以孙光荣著述及选编文献为主，尚未纳入弟子整理的大规模门诊医案。后续将扩充医案数据、开展专家评价，"
      "并探索基于图谱路径的辨证推理与个性化用药推荐。")

    # ---------------- 4 结论 ----------------
    heading(doc, "4  结论", 1)
    P("本研究构建了包含 19 343 个实体、25 713 条关系的孙光荣中和学术思想专题知识图谱，几乎全部关系可回溯至原文并标注知识来源；"
      "图谱清晰呈现了中和思想“经典渊源—核心理念—治则—治法—三联药组”的传承与落实路径，以及以“参—芪—丹”为基础的用药规律。"
      "在此基础上提出的可溯源 GraphRAG 方法在召回率上优于对比方案，并显著改善了检索证据的侧面均衡与证据质量，"
      "为名老中医学术思想的数字化、智能化传承提供了可复用的技术路径。")

    heading(doc, "利益冲突声明", 2)
    para(doc, "【全体作者声明无利益冲突。】", indent=False)
    heading(doc, "数据与代码可获得性", 2)
    para(doc, "图谱数据、问答系统源码、检索评测题集与全部图表生成脚本见【公开仓库地址】。", indent=False)

    # ---------------- 参考文献 ----------------
    heading(doc, "参考文献", 1)
    refs = [
        "刘应科, 孙光荣. 中医临证四大核心理念之中和观[J]. 湖南中医药大学学报, 2016, 36(9): 【页码】.",
        "薛武更. 孙光荣“三联药组”配伍学术特色[J]. 山东中医杂志, 2017, 36(10): 【页码】.",
        "HOGAN A, BLOMQVIST E, COCHEZ M, et al. Knowledge graphs[J]. ACM Computing Surveys, 2021, 54(4): 71.",
        "JI S X, PAN S R, CAMBRIA E, et al. A survey on knowledge graphs: representation, acquisition, and applications[J]. "
        "IEEE Transactions on Neural Networks and Learning Systems, 2022, 33(2): 494-514.",
        "JI Z W, LEE N, FRIESKE R, et al. Survey of hallucination in natural language generation[J]. ACM Computing Surveys, "
        "2023, 55(12): 248.",
        "HUANG L, YU W J, MA W T, et al. A survey on hallucination in large language models: principles, taxonomy, challenges, "
        "and open questions[J]. ACM Transactions on Information Systems, 2025, 43(2): 42.",
        "LEWIS P, PEREZ E, PIKTUS A, et al. Retrieval-augmented generation for knowledge-intensive NLP tasks[C]//Advances in "
        "Neural Information Processing Systems 33. Red Hook: Curran Associates, 2020: 9459-9474.",
        "GAO Y F, XIONG Y, GAO X Y, et al. Retrieval-augmented generation for large language models: a survey[EB/OL]. "
        "(2023-12-18)[【引用日期】]. https://arxiv.org/abs/2312.10997.",
        "EDGE D, TRINH H, CHENG N, et al. From local to global: a graph RAG approach to query-focused summarization[EB/OL]. "
        "(2024-04-24)[【引用日期】]. https://arxiv.org/abs/2404.16130.",
        "PAN S R, LUO L H, WANG Y F, et al. Unifying large language models and knowledge graphs: a roadmap[J]. IEEE Transactions "
        "on Knowledge and Data Engineering, 2024, 36(7): 3580-3599.",
        "ROBERTSON S, ZARAGOZA H. The probabilistic relevance framework: BM25 and beyond[J]. Foundations and Trends in "
        "Information Retrieval, 2009, 3(4): 333-389.",
        "JACOMY M, VENTURINI T, HEYMANN S, et al. ForceAtlas2, a continuous graph layout algorithm for handy network "
        "visualization designed for the Gephi software[J]. PLoS One, 2014, 9(6): e98679.",
        "CLAUSET A, SHALIZI C R, NEWMAN M E J. Power-law distributions in empirical data[J]. SIAM Review, 2009, 51(4): 661-703.",
        "BLONDEL V D, GUILLAUME J L, LAMBIOTTE R, et al. Fast unfolding of communities in large networks[J]. Journal of "
        "Statistical Mechanics: Theory and Experiment, 2008, 2008(10): P10008.",
        "AGRAWAL R, SRIKANT R. Fast algorithms for mining association rules[C]//Proceedings of the 20th International Conference "
        "on Very Large Data Bases. San Francisco: Morgan Kaufmann, 1994: 487-499.",
        "WILSON E B. Probable inference, the law of succession, and statistical inference[J]. Journal of the American "
        "Statistical Association, 1927, 22(158): 209-212.",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.7)
        p.paragraph_format.line_spacing = 1.2
        rich(p, f"[{i}] ".replace("[", "［").replace("]", "］") + ref, 9)
    para(doc, "【投稿前请补充近 3 年中医药知识图谱、名老中医经验传承数字化、中医大模型问答方向的中文核心期刊文献 4—8 篇，"
              "并核对上列文献的页码与引用日期。】", 9, indent=False, before=4)

    # ---------------- 附录 ----------------
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    heading(doc, "附录（补充材料）", 1)
    TABLE_MAP.update({"表2": "附表S1", "表3": "附表S2"})
    tables.table2(doc)
    tables.table3(doc)
    FIG_N[0] = 0
    for png, zh, en in [
        ("图3_实体与关系类型分布.png", "附图S1  实体与关系类型分布", "Fig.S1  Distribution of entity and relation types"),
        ("图5_知识质量评估.png", "附图S2  知识质量评估", "Fig.S2  Quality assessment of the knowledge graph"),
        ("图6_网络拓扑特征.png", "附图S3  网络拓扑特征", "Fig.S3  Topological characteristics of the knowledge graph"),
    ]:
        pp = doc.add_paragraph()
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pp.add_run().add_picture(os.path.join(PNG, png), width=Cm(15.5))
        for s, east in ((zh, "黑体"), (en, "宋体")):
            c = doc.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c.add_run(s)
            set_font(r, 9, east, bold=(east == "黑体"))

    out = os.path.join(OUT, "论文手稿_孙光荣中和学术思想专题知识图谱与可溯源智能问答研究.docx")
    zoom = doc.settings.element.find(qn("w:zoom"))      # python-docx 默认模板缺 percent 属性
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    doc.save(out)
    print("  ✔", out)
    return out


if __name__ == "__main__":
    build()
